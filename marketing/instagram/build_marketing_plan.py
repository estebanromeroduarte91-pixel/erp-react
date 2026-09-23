from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Pixit-Plan-Marketing-Instagram.docx"

BLUE = "245FFF"
DARK = "081328"
MID = "596980"
LIGHT = "EEF3FF"
BORDER = "D9DEE8"
PALE = "F6F8FC"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)
    keep.set(qn("w:val"), "1" if value else "0")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("7B8799")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def set_run_font(run, name="Arial", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    set_run_font(r, size=10, color=DARK)
    return p


def add_label(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text.upper())
    set_run_font(r, size=8.5, bold=True, color=BLUE)
    r.font.all_caps = True
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=10, bold=True, color=DARK)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=10, color=DARK)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10, color=DARK)
    return p


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.76)
sec.left_margin = Inches(0.82)
sec.right_margin = Inches(0.82)
sec.footer_distance = Inches(0.28)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor.from_string(DARK)

title_style = styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
title_style.font.size = Pt(34)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)

if title_style._element.pPr is not None:
    title_border = title_style._element.pPr.find(qn("w:pBdr"))
    if title_border is not None:
        title_style._element.pPr.remove(title_border)

for style_name, size in (("Heading 1", 21), ("Heading 2", 14), ("Heading 3", 11.5)):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(12 if style_name != "Heading 1" else 16)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True

# Footer
footer = sec.footer
footer.is_linked_to_previous = False
fp = footer.paragraphs[0]
fp.clear()
footer_table = footer.add_table(rows=1, cols=2, width=Inches(6.86))
footer_table.columns[0].width = Inches(5.6)
footer_table.columns[1].width = Inches(1.26)
left_p = footer_table.cell(0, 0).paragraphs[0]
left_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
fr = left_p.add_run("PIXIT  |  PLAN DE CONTENIDO PARA INSTAGRAM")
set_run_font(fr, size=8, bold=True, color="7B8799")
right_p = footer_table.cell(0, 1).paragraphs[0]
add_page_number(right_p)
for cell in footer_table.rows[0].cells:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tc_pr.append(borders)
    set_cell_margins(cell, 0, 0, 0, 0)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(76)
r = p.add_run("PIXIT")
set_run_font(r, name="Arial", size=17, bold=True, color=BLUE)

t = doc.add_paragraph(style="Title")
t.alignment = WD_ALIGN_PARAGRAPH.LEFT
t.add_run("Plan de marketing y contenido para Instagram")
title_ppr = t._p.get_or_add_pPr()
title_pborder = title_ppr.find(qn("w:pBdr"))
if title_pborder is not None:
    title_ppr.remove(title_pborder)

sub = doc.add_paragraph()
sub.paragraph_format.space_before = Pt(10)
sub.paragraph_format.space_after = Pt(22)
r = sub.add_run("Estrategia editorial, calendario de publicaciones y captions")
set_run_font(r, size=16, color=MID)

add_body(doc, "Este documento reúne la estrategia acordada para presentar Pixit en Instagram como el sistema de gestión creado para talleres de servicio técnico. La campaña combina contenido comercial, demostraciones del producto y beneficios operativos para atraer dueños y administradores de talleres, generar interés y conducirlos hacia una prueba gratuita o una demostración.")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(80)
r = p.add_run("Formato principal")
set_run_font(r, size=9, bold=True, color=BLUE)
p2 = doc.add_paragraph()
p2.paragraph_format.space_after = Pt(5)
r = p2.add_run("Feed de Instagram 1080 x 1350 px")
set_run_font(r, name="Arial", size=20, bold=True, color=DARK)
p3 = doc.add_paragraph()
r = p3.add_run("Proporción vertical 4:5")
set_run_font(r, size=11, color=MID)

doc.add_page_break()

# Strategy
doc.add_heading("Estrategia general", level=1)
add_body(doc, "La estrategia posiciona a Pixit como una solución especializada para talleres de reparación de celulares y computadores. El contenido debe demostrar cómo el sistema reduce el desorden operativo, centraliza la información y entrega mayor control sobre las órdenes, el equipo de trabajo, el inventario y las finanzas.")

doc.add_heading("Objetivo principal", level=2)
add_body(doc, "Generar reconocimiento y solicitudes de prueba mostrando problemas habituales del taller y la forma concreta en que Pixit los resuelve.")

doc.add_heading("Audiencia", level=2)
add_bullet(doc, "Dueños y administradores de talleres de reparación.")
add_bullet(doc, "Servicios técnicos de celulares, computadores y equipos electrónicos.")
add_bullet(doc, "Negocios que operan con papel, Excel y conversaciones dispersas en WhatsApp.")
add_bullet(doc, "Talleres que necesitan controlar técnicos, sucursales, repuestos y proveedores.")

doc.add_heading("Posicionamiento", level=2)
add_body(doc, "Pixit es el ERP y punto de venta diseñado para servicios técnicos. Conecta órdenes de trabajo, clientes, inventario, ventas, compras, gastos, comisiones, permisos y estadísticas dentro de una misma plataforma.")

doc.add_heading("Propuesta de valor", level=2)
add_body(doc, "Todo lo que ocurre en el taller queda registrado en un solo lugar. El beneficio central no es únicamente digitalizar tareas: es reducir errores, ahorrar tiempo, mejorar la experiencia del cliente y tomar decisiones con información real.")

doc.add_heading("Mensajes centrales", level=2)
add_bullet(doc, "Menos desorden operativo y más control sobre cada reparación.")
add_bullet(doc, "Información conectada entre órdenes, ventas, inventario y contabilidad.")
add_bullet(doc, "Una atención profesional para el cliente desde el ingreso hasta la entrega.")
add_bullet(doc, "Control del equipo mediante cargos, permisos y comisiones.")
add_bullet(doc, "Capacidad de crecer desde un taller individual hasta una operación con varias sucursales.")

doc.add_heading("Tono de comunicación", level=2)
add_body(doc, "Claro, directo y profesional. El contenido debe hablar desde los problemas cotidianos del taller, evitar tecnicismos innecesarios y explicar cada función mediante su beneficio práctico.")

doc.add_page_break()

# Pillars
doc.add_heading("Pilares de contenido", level=1)

pillars = [
    ("Operación del taller", "Órdenes de trabajo, flujo de estados, fotografías mediante QR y derivaciones. Este pilar demuestra orden, trazabilidad y rapidez."),
    ("Ventas e inventario", "Punto de venta, existencias por sucursal y actualización de stock. Muestra cómo Pixit conecta la venta con los movimientos reales del negocio."),
    ("Finanzas y rentabilidad", "Gastos, órdenes de compra, proveedores, resultados y comisiones. Ayuda a comunicar que Pixit permite entender cuánto gana realmente el taller."),
    ("Equipo y seguridad", "Usuarios, cargos, permisos y técnicos responsables. Refuerza el control interno y la posibilidad de delegar con confianza."),
    ("Experiencia del cliente", "Notificaciones, presupuestos y seguimiento. Presenta a Pixit como una herramienta para entregar un servicio más profesional."),
    ("Crecimiento", "Planes y sucursales. Muestra que el sistema acompaña distintas etapas del negocio."),
]

table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Inches(1.85)
table.columns[1].width = Inches(4.75)
hdr = table.rows[0]
set_repeat_table_header(hdr)
for i, text in enumerate(("Pilar", "Enfoque")):
    cell = hdr.cells[i]
    cell.text = text
    set_cell_shading(cell, DARK)
    set_cell_borders(cell)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in cell.paragraphs[0].runs:
        set_run_font(run, size=10, bold=True, color="FFFFFF")

for idx, (name, desc) in enumerate(pillars):
    cells = table.add_row().cells
    for cell in cells:
        set_cell_shading(cell, "FFFFFF" if idx % 2 == 0 else PALE)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cells[0].text = name
    cells[1].text = desc
    for run in cells[0].paragraphs[0].runs:
        set_run_font(run, size=9.5, bold=True, color=DARK)
    for run in cells[1].paragraphs[0].runs:
        set_run_font(run, size=9.5, color=DARK)

doc.add_heading("Dirección visual", level=2)
add_bullet(doc, "Fondo blanco y composición limpia con amplio espacio visual.")
add_bullet(doc, "Titulares breves, grandes y de alto peso tipográfico.")
add_bullet(doc, "Azul eléctrico y degradado azul a violeta como acentos principales.")
add_bullet(doc, "Mockups del producto, tarjetas flotantes y sombras suaves.")
add_bullet(doc, "Una sola idea principal por publicación.")

doc.add_heading("Formatos", level=2)
add_body(doc, "Las piezas principales se publican en el feed con formato vertical 1080 x 1350 px. Cuando una función necesite explicar varios pasos, el contenido puede transformarse en carrusel. Los videos de demostración deben utilizar formato Reel 1080 x 1920 px.")

doc.add_page_break()

# Calendar
doc.add_heading("Calendario sugerido", level=1)
add_body(doc, "La campaña se puede publicar durante cuatro semanas, con dos o tres publicaciones semanales. El orden alterna contenido de marca, operación, finanzas y control interno para mostrar la amplitud del producto sin saturar con funciones similares.")

calendar = [
    ("Semana 1", "Lunes", "Simplifica tu taller", "Marca y propuesta de valor"),
    ("Semana 1", "Miércoles", "Órdenes sin caos", "Operación"),
    ("Semana 1", "Viernes", "Inventario y POS", "Ventas e inventario"),
    ("Semana 2", "Martes", "Clientes informados", "Experiencia del cliente"),
    ("Semana 2", "Jueves", "Comisiones para técnicos", "Equipo y finanzas"),
    ("Semana 3", "Lunes", "Fotografías mediante QR", "Operación"),
    ("Semana 3", "Miércoles", "Accesos según cargo", "Equipo y seguridad"),
    ("Semana 3", "Viernes", "Gastos y órdenes de compra", "Finanzas"),
    ("Semana 4", "Martes", "Derivaciones con registro", "Operación"),
    ("Semana 4", "Jueves", "Planes para crecer", "Conversión"),
]

table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [1.0, 0.9, 2.6, 2.0]
for i, w in enumerate(widths):
    table.columns[i].width = Inches(w)
hdr = table.rows[0]
set_repeat_table_header(hdr)
for i, text in enumerate(("Semana", "Día", "Publicación", "Pilar")):
    cell = hdr.cells[i]
    cell.text = text
    set_cell_shading(cell, DARK)
    set_cell_borders(cell)
    set_cell_margins(cell, 100, 100, 100, 100)
    for run in cell.paragraphs[0].runs:
        set_run_font(run, size=9, bold=True, color="FFFFFF")
for idx, row in enumerate(calendar):
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
        set_cell_shading(cells[i], "FFFFFF" if idx % 2 == 0 else PALE)
        set_cell_borders(cells[i])
        set_cell_margins(cells[i], 90, 100, 90, 100)
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cells[i].paragraphs[0].runs:
            set_run_font(run, size=8.8, bold=(i == 2), color=DARK)

doc.add_heading("Ritmo de publicación", level=2)
add_body(doc, "Publicar entre las 11:00 y las 14:00 o entre las 18:00 y las 21:00, y ajustar el horario según el rendimiento real de la cuenta. Mantener al menos un día entre publicaciones facilita la distribución y permite responder consultas.")

doc.add_heading("Distribución recomendada", level=2)
add_bullet(doc, "Publicación individual para mensajes de marca, planes, comisiones, permisos y finanzas.")
add_bullet(doc, "Carrusel para órdenes, QR y derivaciones cuando se quiera explicar el proceso paso a paso.")
add_bullet(doc, "Reel para demostraciones reales dentro del software.")

doc.add_page_break()

# Posts detailed
posts = [
    {
        "n": 1,
        "title": "Simplifica tu taller",
        "format": "Publicación de marca",
        "cover": "Simplifica tu taller. Multiplica tus ventas.",
        "support": "Órdenes, clientes, inventario y ventas conectados en un solo lugar.",
        "caption": "¿Tu servicio técnico todavía depende de papeles, planillas y mensajes dispersos? Pixit conecta órdenes, clientes, inventario y ventas en una sola plataforma creada para talleres de servicio técnico.",
        "cta": "Prueba Pixit gratis durante 30 días en pixit.cl",
    },
    {
        "n": 2,
        "title": "Órdenes sin caos",
        "format": "Carrusel o publicación individual",
        "cover": "Cada reparación. Siempre bajo control.",
        "support": "Cliente, equipo, falla, técnico y estado dentro de una orden digital.",
        "slides": ["Ingreso", "Chequeo", "Reparación", "Entrega"],
        "caption": "Registra al cliente, el equipo, la falla, el técnico responsable y cada cambio de estado dentro de una orden digital. Tu equipo encuentra la información que necesita sin perder tiempo buscando mensajes.",
        "cta": "Ordena la operación de tu taller con Pixit.",
    },
    {
        "n": 3,
        "title": "Inventario y POS conectados",
        "format": "Publicación individual",
        "cover": "Vende. Descuenta stock. Todo en un paso.",
        "support": "Conecta ventas, repuestos y existencias por sucursal sin duplicar trabajo.",
        "caption": "Vende accesorios o repuestos y actualiza el inventario dentro del mismo flujo. Revisa las existencias por sucursal y evita registrar dos veces el mismo movimiento.",
        "cta": "Vende con más control desde pixit.cl",
    },
    {
        "n": 4,
        "title": "Clientes informados",
        "format": "Carrusel",
        "cover": "Menos preguntas. Más confianza.",
        "support": "Mantén informado a tu cliente desde el ingreso hasta la entrega.",
        "slides": ["Equipo recibido", "Presupuesto disponible", "Equipo listo para retirar", "Historial y seguimiento de la orden"],
        "caption": "Mantén informado al cliente desde el ingreso del equipo hasta el momento del retiro. Pixit centraliza la comunicación y ayuda a entregar una experiencia más profesional.",
        "cta": "Mejora la experiencia de tus clientes con Pixit.",
    },
    {
        "n": 5,
        "title": "Planes para crecer",
        "format": "Publicación comercial",
        "cover": "Empieza simple. Crece sin límites.",
        "support": "Planes Starter, PRO y Scale para acompañar cada etapa del taller.",
        "caption": "Digitaliza tu taller con usuarios ilimitados y elige las funciones que necesita tu operación. Planes Starter, PRO y Scale desde 0,5 UF + IVA al mes.",
        "cta": "Prueba todas las funciones gratis durante 30 días.",
    },
    {
        "n": 6,
        "title": "Comisiones para técnicos",
        "format": "Carrusel o publicación individual",
        "cover": "Calcula comisiones. Sin planillas.",
        "support": "Asigna al técnico, define el porcentaje y registra el pago desde la orden.",
        "slides": ["Asigna el técnico responsable", "Define el monto y el porcentaje", "Pixit calcula la comisión", "Registra el pago", "El pago queda incorporado como gasto"],
        "caption": "Asigna el técnico responsable, define el porcentaje comisionable y deja que Pixit calcule el monto. Cuando registras el pago, la comisión queda marcada como pagada e incorporada como gasto.",
        "cta": "Controla órdenes y comisiones en un solo lugar.",
    },
    {
        "n": 7,
        "title": "Gastos y órdenes de compra",
        "format": "Publicación individual",
        "cover": "No basta con vender. Controla lo que gastas.",
        "support": "Registra gastos, compra a proveedores y conoce el resultado real del negocio.",
        "caption": "Vender más es solo una parte del resultado. Registra gastos, crea órdenes de compra para tus proveedores y actualiza el inventario cuando recibes los productos.",
        "cta": "Conoce la rentabilidad real de tu negocio.",
    },
    {
        "n": 8,
        "title": "Fotografías mediante QR",
        "format": "Reel o carrusel",
        "cover": "Escanea. Captura. Listo.",
        "support": "Sin descargar archivos ni enviarte fotos por WhatsApp.",
        "slides": ["Abre la orden de trabajo", "Genera el código QR", "Escanéalo con el teléfono", "Toma las fotografías", "Las imágenes quedan vinculadas a la orden"],
        "caption": "No necesitas descargar las fotografías ni enviártelas por WhatsApp. Escanea el QR de la orden con tu teléfono, toma las fotos y guárdalas directamente en el ingreso, la inspección, la salida o una derivación.",
        "cta": "Documenta cada reparación con Pixit.",
    },
    {
        "n": 9,
        "title": "Accesos según el cargo",
        "format": "Publicación individual",
        "cover": "Cada persona ve lo que necesita.",
        "support": "Crea cargos y controla el acceso a cada módulo de Pixit.",
        "caption": "Crea cargos y decide qué módulos puede utilizar cada integrante del equipo. Taller, ventas, inventario, compras, estadísticas y configuración: cada persona ve lo necesario para realizar su trabajo.",
        "cta": "Organiza tu equipo y protege la información del negocio.",
    },
    {
        "n": 10,
        "title": "Derivaciones con registro",
        "format": "Carrusel",
        "cover": "Envía la reparación. No pierdas el control.",
        "support": "Deriva una orden a otro taller y conserva estados, motivo y respaldo fotográfico.",
        "slides": ["Selecciona la orden", "Registra el taller o técnico externo", "Indica el motivo", "Agrega fotografías", "Controla su estado hasta el retorno"],
        "caption": "¿Necesitas enviar una reparación a otro taller? Deriva la orden sin perder su trazabilidad. Registra el destinatario, el motivo, el estado y las fotografías hasta que el equipo regrese.",
        "cta": "Mantén el control incluso cuando el equipo sale de tu taller.",
    },
]

doc.add_heading("Contenido y captions", level=1)
add_body(doc, "Cada publicación utiliza una idea central, un texto breve dentro de la pieza y un caption que amplía el beneficio. Los textos siguientes están listos para copiar, adaptar y publicar.")

for idx, post in enumerate(posts):
    if idx > 0 and idx % 2 == 0:
        doc.add_page_break()
    h = doc.add_heading(f"{post['n']}. {post['title']}", level=2)
    set_keep_with_next(h)
    add_label(doc, post["format"])
    add_body(doc, f"Texto principal: {post['cover']}", bold_lead="Texto principal:")
    add_body(doc, f"Texto de apoyo: {post['support']}", bold_lead="Texto de apoyo:")
    if post.get("slides"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run("Secuencia sugerida")
        set_run_font(r, size=10, bold=True, color=DARK)
        for slide in post["slides"]:
            add_bullet(doc, slide)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Caption")
    set_run_font(r, size=10, bold=True, color=BLUE)
    add_body(doc, post["caption"])
    add_body(doc, f"CTA: {post['cta']}", bold_lead="CTA:")

doc.add_page_break()

# Execution
doc.add_heading("Guía de publicación", level=1)
doc.add_heading("Estructura recomendada del caption", level=2)
add_bullet(doc, "Primera línea: problema o pregunta que el dueño del taller reconozca inmediatamente.")
add_bullet(doc, "Desarrollo: explicación breve de la función y su beneficio práctico.")
add_bullet(doc, "Cierre: una llamada a la acción clara, como probar Pixit, solicitar una demostración o visitar pixit.cl.")

doc.add_heading("Hashtags base", level=2)
add_body(doc, "#Pixit #ServicioTécnico #TallerDeReparación #ReparaciónCelulares #SoftwareParaTalleres #OrdenDeTrabajo #Inventario #PuntoDeVenta")
add_body(doc, "Usar entre cinco y ocho hashtags relevantes por publicación. Mantener #Pixit y variar el resto según el tema para evitar repetir exactamente la misma combinación.")

doc.add_heading("Llamadas a la acción", level=2)
add_bullet(doc, "Prueba Pixit gratis durante 30 días.")
add_bullet(doc, "Conoce cómo funcionaría Pixit en tu taller.")
add_bullet(doc, "Solicita una demostración.")
add_bullet(doc, "Descubre más en pixit.cl.")

doc.add_heading("Gestión de comentarios y mensajes", level=2)
add_body(doc, "Responder durante el mismo día las preguntas sobre precios, módulos y prueba gratuita. Cuando una persona describa un problema concreto de su taller, orientar la respuesta hacia la función correspondiente y ofrecer una demostración.")

doc.add_heading("Uso de demostraciones", level=2)
add_body(doc, "Las capturas y grabaciones deben utilizar datos de ejemplo y evitar nombres, teléfonos, correos, RUT o cifras internas de clientes reales. En los Reels, mostrar acciones breves y visibles: crear una orden, escanear un QR, cambiar un estado o registrar una venta.")

doc.add_heading("Promoción pagada", level=2)
add_body(doc, "Promocionar primero las publicaciones que consigan mejor respuesta orgánica. Priorizar las piezas centradas en problemas fáciles de reconocer, como el desorden de las órdenes, las fotos por WhatsApp, el stock y el cálculo manual de comisiones.")

doc.add_page_break()

# Metrics
doc.add_heading("Medición y optimización", level=1)
add_body(doc, "La campaña debe evaluarse por su capacidad de generar interés comercial y no únicamente por la cantidad de Me gusta. Revisar los resultados semanalmente y comparar cada tema con el promedio de la cuenta.")

metrics = [
    ("Alcance", "Cuántas cuentas vieron la publicación", "Detecta qué temas amplían el reconocimiento"),
    ("Guardados", "Personas que conservaron la publicación", "Señala contenido práctico o educativo"),
    ("Compartidos", "Veces que la pieza fue enviada", "Mide relevancia dentro del rubro"),
    ("Visitas al perfil", "Usuarios que revisaron la cuenta", "Mide interés posterior al contenido"),
    ("Clics al sitio", "Visitas enviadas a pixit.cl", "Conecta el contenido con intención comercial"),
    ("Consultas y pruebas", "Mensajes, demostraciones y registros", "Es el indicador principal de conversión"),
]

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
for i, w in enumerate((1.3, 2.25, 2.95)):
    table.columns[i].width = Inches(w)
hdr = table.rows[0]
set_repeat_table_header(hdr)
for i, text in enumerate(("Indicador", "Qué mide", "Cómo interpretarlo")):
    cell = hdr.cells[i]
    cell.text = text
    set_cell_shading(cell, DARK)
    set_cell_borders(cell)
    set_cell_margins(cell)
    for run in cell.paragraphs[0].runs:
        set_run_font(run, size=9.2, bold=True, color="FFFFFF")
for idx, row in enumerate(metrics):
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
        set_cell_shading(cells[i], "FFFFFF" if idx % 2 == 0 else PALE)
        set_cell_borders(cells[i])
        set_cell_margins(cells[i])
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cells[i].paragraphs[0].runs:
            set_run_font(run, size=9.2, bold=(i == 0), color=DARK)

doc.add_heading("Criterios para ajustar la campaña", level=2)
add_bullet(doc, "Repetir los temas que generen más guardados, compartidos y consultas.")
add_bullet(doc, "Transformar las publicaciones con mayor interés en Reels demostrativos.")
add_bullet(doc, "Probar dos versiones del titular cuando una función sea valiosa pero tenga poco alcance.")
add_bullet(doc, "Mantener los mensajes comerciales dentro de una mezcla equilibrada con contenido educativo y demostrativo.")

doc.add_heading("Siguiente etapa", level=2)
add_body(doc, "Después de publicar las diez piezas, crear una segunda campaña basada en los resultados. Las funciones con mejor respuesta pueden desarrollarse como tutoriales breves, casos de uso, preguntas frecuentes o testimonios de talleres.")

doc.core_properties.title = "Plan de marketing y contenido para Instagram de Pixit"
doc.core_properties.subject = "Estrategia, calendario y captions para Instagram"
doc.core_properties.author = "Pixit"
doc.core_properties.keywords = "Pixit, Instagram, marketing, contenido, captions"

doc.save(OUT)
print(OUT)
